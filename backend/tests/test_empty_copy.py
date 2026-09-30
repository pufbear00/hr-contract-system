"""Run only on a temporary copy; never create data in the distributable tree."""
import asyncio
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile


def test_isolated_empty_startup_and_local_workflow():
    backend = Path(__file__).resolve().parents[1]
    with tempfile.TemporaryDirectory(prefix="hr-empty-test-") as temporary:
        isolated = Path(temporary) / "backend"
        shutil.copytree(backend, isolated, ignore=shutil.ignore_patterns(
            "local_runtime", ".venv", "__pycache__", ".pytest_cache"))
        script = r'''
import asyncio
import io
import json
import secrets
import socket
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from sqlalchemy import text
from app.config import settings, BASE_DIR
from app.database import engine, async_session, Base
from app.main import app
from init_local import create_admin

def forbid(*args, **kwargs):
    raise AssertionError("No external network is allowed in an empty copy")

async def counts():
    async with async_session() as session:
        return {table.name: (await session.execute(text('SELECT COUNT(*) FROM "' + table.name + '"'))).scalar_one()
                for table in Base.metadata.sorted_tables}

with patch.object(socket.socket, "connect", forbid):
    assert settings.OFFLINE_MODE
    assert settings.DATABASE_URL == f"sqlite+aiosqlite:///{BASE_DIR / 'local_runtime' / 'empty_hr.db'}"
    assert settings.UPLOAD_DIR.resolve().is_relative_to(BASE_DIR)
    assert settings.GENERATED_DIR.resolve().is_relative_to(BASE_DIR)
    with TestClient(app) as client:
        assert client.get("/api/v1/health").json()["offline"] is True
        assert all(value == 0 for value in asyncio.run(counts()).values())
        password = secrets.token_urlsafe(24)
        asyncio.run(create_admin("synthetic-demo-user", password))
        initial = asyncio.run(counts())
        assert initial.pop("users") == 1
        assert all(value == 0 for value in initial.values())
        login = client.post("/api/v1/auth/login", json={"username": "synthetic-demo-user", "password": password})
        assert login.status_code == 200, login.text
        client.headers["Authorization"] = "Bearer " + login.json()["access_token"]
        assert client.get("/api/v1/employees").json() == []
        assert client.get("/api/v1/companies").json() == []
        assert all(value == 0 for value in client.get("/api/v1/data/stats").json().values())
        assert client.post("/api/v1/employees/feishu-sync").status_code == 503
        assert client.post("/api/v1/data/import", files={"file": ("synthetic.zip", b"empty")}).status_code == 503
        created = client.post("/api/v1/employees", json={"name": "虚构演示人员", "employee_type": "演示"})
        assert created.status_code == 200, created.text
        assert client.get("/api/v1/employees").json()[0]["name"] == "虚构演示人员"
        assert client.delete("/api/v1/employees/" + str(created.json()["id"])).status_code == 200
        assert client.get("/api/v1/employees").json() == []
        assert client.get("/api/v1/data/export").status_code == 200
        from docx import Document
        from app.services.doc_generator import generate_zip_for_employee
        synthetic_template = Document()
        synthetic_template.add_paragraph("虚构演示：{{ employee_name }}")
        template_buffer = io.BytesIO()
        synthetic_template.save(template_buffer)
        uploaded = client.post("/api/v1/templates/upload", data={"name": "虚构演示模板"},
                               files={"file": ("../synthetic.docx", template_buffer.getvalue())})
        assert uploaded.status_code == 200, uploaded.text
        assert uploaded.json()["original_filename"] == "synthetic.docx"
        template_path = next(settings.UPLOAD_DIR.glob("*.docx"))
        generated = generate_zip_for_employee("../../虚构演示人员",
            [{"name": "../../synthetic", "path": str(template_path)}],
            {"employee_name": "虚构演示人员"}, settings.GENERATED_DIR / "synthetic-job")
        assert len(generated) == 1 and generated[0].resolve().is_relative_to(settings.GENERATED_DIR)
        assert "虚构演示人员" in Document(generated[0]).paragraphs[0].text
print(json.dumps({"empty_business_tables": True, "fresh_admin": True, "local_crud": True,
                  "live_sync_blocked": True, "backup_import_blocked": True, "network_blocked": True,
                  "synthetic_template_generation": True, "local_paths_confined": True}))
'''
        environment = os.environ.copy()
        environment.update({
            "HR_DATABASE_URL": "sqlite+aiosqlite:///DO_NOT_USE_ORIGINAL.db",
            "HR_UPLOAD_DIR": "/DO_NOT_USE_ORIGINAL/uploads",
            "HR_FEISHU_APP_SECRET": "synthetic-ignored-marker",
            "PYTHONDONTWRITEBYTECODE": "1",
        })
        result = subprocess.run([sys.executable, "-c", script], cwd=isolated,
                                env=environment, text=True, capture_output=True)
        assert result.returncode == 0, result.stderr
        assert json.loads(result.stdout.strip())["empty_business_tables"]
