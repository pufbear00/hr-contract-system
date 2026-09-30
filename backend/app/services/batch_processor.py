import io
from openpyxl import load_workbook


def parse_excel(file_bytes: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """Parse an Excel file and return (headers, rows).

    Each row is a dict mapping header → cell value.
    """
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h).strip() for h in next(rows_iter) if h is not None]

    rows = []
    for row_values in rows_iter:
        values = list(row_values)[:len(headers)]
        if all(v is None for v in values):
            continue
        row_dict = {}
        for i, header in enumerate(headers):
            val = values[i] if i < len(values) else None
            row_dict[header] = str(val).strip() if val is not None else ""
        rows.append(row_dict)

    wb.close()
    return headers, rows
