import io
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.database import get_db
from app.auth import get_current_user
from app.models.user import User
from app.models.employee import Employee
from app.models.template import ContractTemplate
from app.models.template_set import TemplateSet, TemplateSetItem
from app.models.generation import GenerationJob, GeneratedDocument
from app.models.social_wage import SocialAverageWage
from app.schemas.employee import (
    EmployeeCreate, EmployeeUpdate, EmployeeResponse, EmployeeGenerateRequest,
)
from app.services.doc_generator import generate_zip_for_employee, create_zip
from app.services.feishu_sync import sync_employees
from app.api.companies import get_company_values
from app.config import settings

router = APIRouter(prefix="/api/v1/employees", tags=["employees"])

# 员工固定字段 → 模板变量的映射
EMPLOYEE_TO_TEMPLATE = {
    "name": "employee_name",
    "id_number": "id_number",
    "gender": "gender",
    "phone": "phone",
    "address": "address",
    "position": "position",
    "department": "department",
    "entry_date": "entry_date",
}

# data 中文 key → 模板变量 key 的映射（解决 key 不一致问题）
DATA_TO_TEMPLATE_VAR = {
    "正式员工月薪": "salary",
    "实习生日薪": "daily_pay",
    "租房补贴": "housing_allowance",
}


def _employee_to_template_values(emp: Employee) -> dict[str, str]:
    """从员工档案提取所有可用的模板变量值。"""
    values = {}
    # 固定字段映射
    for field, tpl_var in EMPLOYEE_TO_TEMPLATE.items():
        val = getattr(emp, field, None)
        if val is not None and val != "":
            values[tpl_var] = str(val)
    # data JSON 中的 key 直接对应模板变量名
    if emp.data:
        for k, v in emp.data.items():
            if v is not None and v != "":
                values[k] = str(v)

    # 中文 key → 模板变量 key 翻译
    for data_key, tpl_key in DATA_TO_TEMPLATE_VAR.items():
        if tpl_key not in values and data_key in values:
            values[tpl_key] = values[data_key]

    # 学校 + 专业 → school 模板变量（"学校--专业" 格式）
    if "school" not in values:
        school_name = values.get("school_name", "")
        major = values.get("major", "")
        if school_name and major:
            values["school"] = f"{school_name}--{major}"
        elif school_name:
            values["school"] = school_name
    elif "major" in values:
        # school 已有但 major 单独存储，拼合
        school_val = values["school"]
        if "--" not in school_val:
            values["school"] = f"{school_val}--{values['major']}"

    # position → intern_position fallback（实习合同用）
    if "intern_position" not in values and "position" in values:
        values["intern_position"] = values["position"]

    # 自动推导 base_salary（用于 N+1 补偿计算）
    if "base_salary" not in values and emp.data:
        salary = emp.data.get("正式员工月薪") or emp.data.get("实习生月薪")
        if salary:
            values["base_salary"] = str(salary)
    return values


# ─── Batch Import ───────────────────────────────────────

EXCEL_TO_FIELD = {
    "姓名": "name",
    "员工类型": "employee_type",
    "身份证号": "id_number",
    "性别": "gender",
    "手机": "phone",
    "地址": "address",
    "职位": "position",
    "部门": "department",
    "入职日期": "entry_date",
    "状态": "status",
}

# Excel 列名 → data JSON key（用于存入模板兼容的 key）
EXCEL_TO_DATA_KEY = {
    "学校": "school_name",
    "专业": "major",
    "租房补贴": "housing_allowance",
}

TEMPLATE_HEADERS = list(EXCEL_TO_FIELD.keys()) + [
    "正式员工月薪", "实习生日薪", "实习生月薪", "学校", "专业", "租房补贴",
]


@router.get("/template/download")
async def download_template(user: User = Depends(get_current_user)):
    """下载员工导入 Excel 模板。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "员工导入模板"
    ws.append(TEMPLATE_HEADERS)

    col_widths = [12, 12, 22, 8, 15, 30, 15, 15, 15, 10, 15, 12, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=employee_template.xlsx"},
    )


# 固定字段 → Excel 列名（EXCEL_TO_FIELD 的反转）
FIELD_TO_EXCEL = {v: k for k, v in EXCEL_TO_FIELD.items()}
# data key → Excel 列名（EXCEL_TO_DATA_KEY 的反转）
DATA_KEY_TO_EXCEL = {v: k for k, v in EXCEL_TO_DATA_KEY.items()}
# data 中直接用中文 key 的列
DATA_CHINESE_KEYS = ["正式员工月薪", "实习生日薪", "实习生月薪"]


@router.get("/export")
async def export_employees(
    search: str | None = Query(None),
    employee_type: str | None = Query(None),
    status: str | None = Query(None),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """导出员工档案为 Excel。"""
    query = select(Employee).order_by(Employee.entry_date.desc(), Employee.id.desc())
    if search:
        query = query.where(
            Employee.name.contains(search)
            | Employee.phone.contains(search)
            | Employee.id_number.contains(search)
        )
    if employee_type:
        query = query.where(Employee.employee_type == employee_type)
    if status:
        query = query.where(Employee.status == status)
    result = await db.execute(query)
    employees = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "员工档案"
    ws.append(TEMPLATE_HEADERS)

    for emp in employees:
        row = []
        for header in TEMPLATE_HEADERS:
            field = EXCEL_TO_FIELD.get(header)
            if field:
                row.append(getattr(emp, field, None) or "")
            else:
                # 从 data JSON 取值
                data_key = EXCEL_TO_DATA_KEY.get(header)
                if data_key and emp.data:
                    row.append(emp.data.get(data_key, "") or "")
                elif header in DATA_CHINESE_KEYS and emp.data:
                    row.append(emp.data.get(header, "") or "")
                else:
                    row.append("")
        ws.append(row)

    col_widths = [12, 12, 22, 8, 15, 30, 15, 15, 15, 10, 15, 12, 12, 12, 12]
    for i, w in enumerate(col_widths, 1):
        if i <= ws.max_column:
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=employees_export.xlsx"},
    )


@router.post("/batch-upload")
async def batch_upload(
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
):
    """上传 Excel → 返回预览数据（headers + rows）。"""
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="请上传 .xlsx 或 .xls 文件")

    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(c).strip() if c else "" for c in next(rows_iter, [])]

    if not headers:
        raise HTTPException(status_code=400, detail="Excel 文件为空")

    rows = []
    for row in rows_iter:
        row_dict = {}
        for i, val in enumerate(row):
            if i < len(headers) and headers[i]:
                if isinstance(val, datetime):
                    val = val.date().isoformat()
                elif isinstance(val, date):
                    val = val.isoformat()
                elif val is not None:
                    val = str(val).strip()
                row_dict[headers[i]] = val
        if any(v for v in row_dict.values()):
            rows.append(row_dict)

    wb.close()

    return {"headers": headers, "rows": rows, "total": len(rows)}


@router.post("/batch", response_model=list[EmployeeResponse])
async def batch_create(
    rows: list[dict],
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """确认批量创建员工记录。"""
    created = []
    for row in rows:
        emp_data: dict = {}
        extra_data: dict = {}

        for col_name, val in row.items():
            field = EXCEL_TO_FIELD.get(col_name)
            if field:
                emp_data[field] = val if val else None
            elif col_name and val:
                # 中文列名转为模板兼容的 data key
                data_key = EXCEL_TO_DATA_KEY.get(col_name, col_name)
                extra_data[data_key] = val

        if not emp_data.get("name"):
            continue

        emp_data.setdefault("employee_type", "正式")
        emp_data.setdefault("status", "在职")
        emp_data["data"] = extra_data if extra_data else {}

        emp = Employee(**emp_data)
        db.add(emp)
        created.append(emp)

    await db.commit()
    for emp in created:
        await db.refresh(emp)
    return created


# ─── 飞书同步 ────────────────────────────────────────────

@router.post("/feishu-sync")
async def feishu_sync(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """一键同步飞书人事花名册。"""
    result = await sync_employees(db)
    return result


# ─── CRUD ───────────────────────────────────────────────

@router.post("", response_model=EmployeeResponse)
async def create_employee(
    body: EmployeeCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    emp = Employee(**body.model_dump())
    db.add(emp)
    await db.commit()
    await db.refresh(emp)
    return emp


@router.get("", response_model=list[EmployeeResponse])
async def list_employees(
    search: str | None = Query(None),
    employee_type: str | None = Query(None),
    status: str | None = Query(None),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    query = select(Employee).order_by(Employee.entry_date.desc(), Employee.id.desc())
    if search:
        query = query.where(
            Employee.name.contains(search)
            | Employee.phone.contains(search)
            | Employee.id_number.contains(search)
        )
    if employee_type:
        query = query.where(Employee.employee_type == employee_type)
    if status:
        query = query.where(Employee.status == status)
    result = await db.execute(query)
    return result.scalars().all()


@router.get("/{employee_id}", response_model=EmployeeResponse)
async def get_employee(
    employee_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    result = await db.execute(select(Employee).where(Employee.id == employee_id))
    emp = result.scalar_one_or_none()
    if not emp:
        raise HTTPException(status_code=404, detail="员工不存在")
    return emp


@router.put("/{employee_id}", response_model=EmployeeResponse)
async def update_employee(
    employee_id: int,
    body: EmployeeUpdate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    result = await db.execute(select(Employee).where(Employee.id == employee_id))
    emp = result.scalar_one_or_none()
    if not emp:
        raise HTTPException(status_code=404, detail="员工不存在")

    update_data = body.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(emp, key, value)
    await db.commit()
    await db.refresh(emp)
    return emp


@router.delete("/{employee_id}")
async def delete_employee(
    employee_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    result = await db.execute(select(Employee).where(Employee.id == employee_id))
    emp = result.scalar_one_or_none()
    if not emp:
        raise HTTPException(status_code=404, detail="员工不存在")
    await db.delete(emp)
    await db.commit()
    return {"message": "删除成功"}


@router.post("/{employee_id}/generate")
async def generate_from_employee(
    employee_id: int,
    body: EmployeeGenerateRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    # 1. 获取员工
    result = await db.execute(select(Employee).where(Employee.id == employee_id))
    emp = result.scalar_one_or_none()
    if not emp:
        raise HTTPException(status_code=404, detail="员工不存在")

    # 2. 获取模板集合
    result = await db.execute(
        select(TemplateSet)
        .options(selectinload(TemplateSet.items).selectinload(TemplateSetItem.template))
        .where(TemplateSet.id == body.template_set_id)
    )
    ts = result.scalar_one_or_none()
    if not ts:
        raise HTTPException(status_code=404, detail="模板集合不存在")

    templates = [item.template for item in ts.items if item.template]
    if not templates:
        raise HTTPException(status_code=400, detail="模板集合中没有模板")

    # 3. 查询最新社平工资上限
    wage_result = await db.execute(
        select(SocialAverageWage).order_by(SocialAverageWage.year.desc()).limit(1)
    )
    wage_row = wage_result.scalar_one_or_none()
    monthly_cap = round(wage_row.annual_amount / 12 * 3, 2) if wage_row else None

    # 4. 合并值：公司值 → 员工已有值 → extra_values（后者覆盖前者）
    merged_values = await get_company_values(body.company_id, db)
    merged_values.update(_employee_to_template_values(emp))
    merged_values.update(body.extra_values)

    employee_name = merged_values.get("employee_name", emp.name)

    # 5. 创建生成任务
    job = GenerationJob(
        job_type="single",
        status="processing",
        template_set_id=ts.id,
        employee_type=ts.employee_type,
        scenario=ts.scenario,
        employee_count=1,
        created_by=user.id,
    )
    db.add(job)
    await db.flush()

    # 5. 生成文档
    job_dir = settings.GENERATED_DIR / f"job_{job.id}"
    tpl_list = [{"name": t.name, "path": t.file_path} for t in templates]

    try:
        saved_files = generate_zip_for_employee(employee_name, tpl_list, merged_values, job_dir, monthly_cap=monthly_cap)
    except Exception as e:
        job.status = "failed"
        job.error_message = str(e)
        await db.commit()
        raise HTTPException(status_code=500, detail=f"生成失败: {str(e)}")

    # 6. 创建 ZIP
    zip_path = settings.GENERATED_DIR / f"job_{job.id}.zip"
    create_zip(job_dir, zip_path)

    job.status = "completed"
    job.zip_file_path = str(zip_path)
    job.document_count = len(saved_files)

    for sf in saved_files:
        doc = GeneratedDocument(
            job_id=job.id,
            employee_name=employee_name,
            template_name=sf.stem,
            filled_values=merged_values,
            file_path=str(sf),
        )
        db.add(doc)

    # 7. 可选：回写新字段到员工 data
    if body.save_back and body.extra_values:
        current_data = dict(emp.data) if emp.data else {}
        current_data.update(body.extra_values)
        emp.data = current_data

    await db.commit()
    return {"job_id": job.id, "message": "生成成功", "document_count": len(saved_files)}
